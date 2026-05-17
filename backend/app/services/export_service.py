import csv
import io
from datetime import datetime, timezone
from typing import List, Dict, Any, Optional
from uuid import UUID

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi.responses import StreamingResponse

from app.core.database import get_supabase


def _fetch_responses_flat(form_id: str) -> tuple[List[str], List[Dict[str, Any]]]:
    db = get_supabase()

    questions = (
        db.table("questions")
        .select("id, label, order_index")
        .eq("form_id", form_id)
        .order("order_index")
        .execute()
        .data
    )

    submissions = (
        db.table("submissions")
        .select("id, submitted_at, last_edited_at, edit_count, user_id")
        .eq("form_id", form_id)
        .execute()
        .data
    )

    form_row = db.table("forms").select("title").eq("id", form_id).single().execute().data
    form_title = form_row["title"] if form_row else ""

    q_labels = [q["label"] for q in questions]
    q_ids = [q["id"] for q in questions]

    base_headers = [
        "submission_id", "submitted_at", "last_edited_at", "edit_count",
        "user_dni", "user_full_name", "user_email", "form_title",
    ]
    headers = base_headers + q_labels

    rows = []
    for sub in submissions:
        user = db.table("users").select("dni, full_name, primary_email, emails").eq("id", sub["user_id"]).single().execute().data or {}
        email = user.get("primary_email") or (user.get("emails") or [""])[0]

        answers_raw = (
            db.table("responses")
            .select("question_id, value")
            .eq("submission_id", sub["id"])
            .execute()
            .data
        )
        ans_map = {a["question_id"]: a["value"] for a in answers_raw}

        row = {
            "submission_id": sub["id"],
            "submitted_at": sub["submitted_at"],
            "last_edited_at": sub.get("last_edited_at", ""),
            "edit_count": sub["edit_count"],
            "user_dni": user.get("dni", ""),
            "user_full_name": user.get("full_name", ""),
            "user_email": email,
            "form_title": form_title,
        }
        for qid, qlabel in zip(q_ids, q_labels):
            row[qlabel] = ans_map.get(qid, "")
        rows.append(row)

    return headers, rows


def export_responses_csv(form_id: str) -> StreamingResponse:
    headers, rows = _fetch_responses_flat(form_id)
    output = io.StringIO()
    output.write("﻿")  # UTF-8 BOM for Excel
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    writer.writerows(rows)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": f'attachment; filename="responses_{form_id}.csv"'},
    )


def export_responses_xlsx(form_id: str) -> StreamingResponse:
    headers, rows = _fetch_responses_flat(form_id)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Responses"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, 2):
        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=row_idx, column=col_idx, value=str(row.get(h, "")))

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        iter([stream.read()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="responses_{form_id}.xlsx"'},
    )


def export_summary_xlsx() -> StreamingResponse:
    db = get_supabase()
    users = db.table("users").select("id, dni, full_name").eq("role", "participant").execute().data
    forms = db.table("forms").select("id, title").execute().data

    submissions = db.table("submissions").select("user_id, form_id").execute().data
    submitted_set = {(s["user_id"], s["form_id"]) for s in submissions}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Completion Summary"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

    headers = ["DNI", "Nombre"] + [f["title"] for f in forms]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill

    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    for row_idx, user in enumerate(users, 2):
        ws.cell(row=row_idx, column=1, value=user.get("dni", ""))
        ws.cell(row=row_idx, column=2, value=user.get("full_name", ""))
        for col_idx, form in enumerate(forms, 3):
            done = (user["id"], form["id"]) in submitted_set
            cell = ws.cell(row=row_idx, column=col_idx, value="SI" if done else "NO")
            cell.fill = green if done else red

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        iter([stream.read()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="completion_summary.xlsx"'},
    )


def export_incomplete_csv() -> StreamingResponse:
    from app.services.report_service import get_incomplete_users
    users = get_incomplete_users()
    headers = ["id", "dni", "full_name", "missing_name", "missing_dni", "no_email", "multiple_emails_no_primary", "pending_forms"]
    output = io.StringIO()
    output.write("﻿")
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for u in users:
        writer.writerow({
            "id": u["id"],
            "dni": u.get("dni", ""),
            "full_name": u.get("full_name", ""),
            "missing_name": u["missing_name"],
            "missing_dni": u["missing_dni"],
            "no_email": u["no_email"],
            "multiple_emails_no_primary": u["multiple_emails_no_primary"],
            "pending_forms": "; ".join(u["pending_forms"]),
        })
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": 'attachment; filename="incomplete_users.csv"'},
    )
